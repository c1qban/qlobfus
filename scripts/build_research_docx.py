from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_inches: float) -> None:
    cell.width = Inches(width_inches)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


INLINE_RE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")


def add_inline_runs(paragraph, text: str) -> None:
    text = text.replace("\\(", "(").replace("\\)", ")").replace("\\[", "[").replace("\\]", "]")
    cursor = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def is_table_separator(line: str) -> bool:
    stripped = line.strip()
    if not stripped.startswith("|") or not stripped.endswith("|"):
        return False
    cells = [cell.strip() for cell in stripped.strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell or "") for cell in cells)


def parse_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def apply_document_styles(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = document.styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    title.font.size = Pt(16)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)
    title.paragraph_format.space_after = Pt(12)

    for style_name, size, color, before, after in [
        ("Heading 1", 16, RGBColor(0x2E, 0x74, 0xB5), 16, 8),
        ("Heading 2", 13, RGBColor(0x2E, 0x74, 0xB5), 12, 6),
        ("Heading 3", 12, RGBColor(0x1F, 0x4D, 0x78), 8, 4),
    ]:
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ["List Bullet", "List Number"]:
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.167


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    col_count = max(len(row) for row in rows)
    rows = [row + [""] * (col_count - len(row)) for row in rows]
    weights: list[int] = []
    for col_idx in range(col_count):
        max_len = max(len(re.sub(r"[*`]", "", row[col_idx])) for row in rows)
        weights.append(max(8, min(max_len, 44)))
    total_weight = sum(weights) or col_count
    widths = [max(0.7, 6.5 * weight / total_weight) for weight in weights]
    width_delta = 6.5 - sum(widths)
    widths[-1] += width_delta

    table = document.add_table(rows=len(rows), cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    table.autofit = False
    set_repeat_table_header(table.rows[0])

    for r_idx, row in enumerate(rows):
        for c_idx, value in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_width(cell, widths[c_idx])
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(2)
            add_inline_runs(paragraph, value)
            for run in paragraph.runs:
                run.font.name = "Calibri"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
                run.font.size = Pt(9.5)
                if r_idx == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor(0x0B, 0x25, 0x45)
            if r_idx == 0:
                set_cell_shading(cell, "F2F4F7")
    document.add_paragraph()


def add_paragraph(document: Document, text: str, style: str | None = None, *, center: bool = False) -> None:
    paragraph = document.add_paragraph(style=style)
    add_inline_runs(paragraph, text)
    if center:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def markdown_to_docx(markdown: str, document: Document) -> None:
    lines = markdown.splitlines()
    in_code = False
    code_lines: list[str] = []
    table_rows: list[list[str]] | None = None
    skip_next_separator = False

    def flush_table() -> None:
        nonlocal table_rows
        if table_rows:
            add_table(document, table_rows)
            table_rows = None

    for idx, raw_line in enumerate(lines):
        line = raw_line.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_table()
            if not in_code:
                in_code = True
                code_lines = []
            else:
                for code_line in code_lines or [""]:
                    paragraph = document.add_paragraph(style="No Spacing")
                    run = paragraph.add_run(code_line)
                    run.font.name = "Consolas"
                    run.font.size = Pt(9.5)
                in_code = False
                code_lines = []
            continue

        if in_code:
            code_lines.append(line)
            continue

        if skip_next_separator:
            skip_next_separator = False
            if is_table_separator(line):
                continue

        if stripped.startswith("|") and stripped.endswith("|"):
            if idx + 1 < len(lines) and is_table_separator(lines[idx + 1]):
                flush_table()
                table_rows = [parse_table_row(line)]
                skip_next_separator = True
                continue
            if table_rows is not None:
                table_rows.append(parse_table_row(line))
                continue
        else:
            flush_table()

        if not stripped:
            continue

        if stripped.startswith(">"):
            add_paragraph(document, stripped.lstrip(">").strip(), "Intense Quote")
            continue

        heading = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if heading:
            level = len(heading.group(1))
            text = heading.group(2).strip()
            if level == 1:
                add_paragraph(document, text, "Title")
            elif level == 2:
                add_paragraph(document, text, "Heading 1")
            else:
                add_paragraph(document, text, "Heading 2")
            continue

        if re.match(r"^\d+\.\s+", stripped):
            add_paragraph(document, re.sub(r"^\d+\.\s+", "", stripped), "List Number")
            continue

        if stripped.startswith("- "):
            add_paragraph(document, stripped[2:], "List Bullet")
            continue

        if stripped in {r"\[", r"\]"}:
            continue

        if stripped.startswith("\\") or stripped.startswith("+") or stripped.startswith("="):
            add_paragraph(document, stripped.strip("\\"), center=True)
            continue

        add_paragraph(document, stripped)

    flush_table()


def build_docx(markdown_path: Path, output_path: Path) -> None:
    document = Document()
    apply_document_styles(document)
    markdown_to_docx(markdown_path.read_text(encoding="utf-8"), document)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser(description="Build research DOCX files from Markdown drafts.")
    parser.add_argument("input", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--title", default="", help="Accepted for backward compatibility; title is read from Markdown.")
    args = parser.parse_args()
    build_docx(args.input, args.output)
    print(args.output)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
